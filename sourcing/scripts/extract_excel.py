#!/usr/bin/env python3
"""Read weekly financing Excel files once; preserve source rows for news screening."""

import argparse
import json
import sys
from datetime import date, datetime
from pathlib import Path


FIELDS = (
    "公司名称", "项目名称", "行业分类", "业务", "最新轮次", "投资时间", "投资金额", "投资方", "url"
)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def header_at(rows):
    for index, row in enumerate(rows[:10]):
        names = [clean(value) for value in row]
        if ("公司名称" in names or "项目名称" in names) and "业务" in names:
            return index, names
    return None


def extract(path):
    suffix = path.suffix.lower()
    if suffix not in {".xlsx", ".xlsm", ".xls"}:
        raise ValueError(f"仅接受 Excel 文件：{path.name}")
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), on_demand=True)
        names = book.sheet_names()

        def read_sheet(name):
            sheet = book.sheet_by_name(name)
            values = []
            for row in sheet.get_rows():
                values.append([
                    xlrd.xldate_as_datetime(c.value, book.datemode)
                    if c.ctype == xlrd.XL_CELL_DATE else c.value for c in row
                ])
            return values

        close = book.release_resources
    else:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        names = book.sheetnames

        def read_sheet(name):
            return list(book[name].iter_rows(values_only=True))

        close = book.close

    try:
        ordered = (["修改后数据"] if "修改后数据" in names else []) + [
            n for n in names if n != "修改后数据"
        ]
        for name in ordered:
            values = read_sheet(name)
            header = header_at(values)
            if header is None:
                continue
            index, labels = header
            positions = {label: i for i, label in enumerate(labels) if label in FIELDS}
            records = []
            for row_number, row in enumerate(values[index + 1:], start=index + 2):
                data = {
                    key: clean(row[i]) if i < len(row) else ""
                    for key, i in positions.items()
                }
                if not (data.get("公司名称") or data.get("项目名称")):
                    continue
                records.append({"source_file": path.name, "sheet": name,
                                "row": row_number, **data})
            return records
        raise ValueError(f"未找到含公司名称/项目名称及业务表头的主表：{path.name}")
    finally:
        close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    if args.output.suffix.lower() != ".json":
        parser.error("输出必须为临时 .json，不能覆盖 Excel")
    try:
        paths = list(dict.fromkeys(p.resolve() for p in args.inputs))
        rows = [row for path in paths for row in extract(path)]
        payload = {"source_files": [p.name for p in paths],
                   "row_count": len(rows), "rows": rows}
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2),
                               encoding="utf-8")
    except (ValueError, OSError, ImportError) as exc:
        print(f"抽取失败：{exc}", file=sys.stderr)
        return 1
    print(f"已读取 {len(paths)} 个 Excel，{len(rows)} 条公司记录。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
